"""Create a read-only inventory of the local 六爻 source collection.

Only filenames, file signatures, hashes, and small text-extraction samples are
examined. Original source files are never modified.
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
from collections import Counter, defaultdict
from pathlib import Path


CORE_TERMS = (
    "六爻", "卜易", "卜筮", "摇钱", "卦例", "断卦", "取用神", "爻位", "六亲", "六神",
)
ADJACENT_TERMS = (
    "八字", "四柱", "风水", "玄空", "面相", "艾灸", "中医", "记忆", "心理暗示", "六壬",
    "阳宅", "阴宅",
)
FIELDS = (
    "relative_path", "extension", "size_bytes", "sha256", "scope_candidate",
    "scope_basis", "format_probe", "sample_detail", "duplicate_of",
)


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def scope_from_filename(path: Path) -> str:
    # Title-only triage. This is never treated as a verified school or doctrine.
    title = path.stem
    core = any(term in title for term in CORE_TERMS)
    adjacent = any(term in title for term in ADJACENT_TERMS)
    if core and adjacent:
        return "mixed_title"
    if core:
        return "liuyao_title"
    if adjacent:
        return "adjacent_title"
    return "review_title"


def probe_pdf(path: Path) -> tuple[str, str]:
    try:
        from pypdf import PdfReader

        reader = PdfReader(str(path), strict=False)
        count = len(reader.pages)
        if not count:
            return "pdf_empty", "pages=0"
        sample = sorted({0, count // 2, count - 1})
        chars = []
        for page_index in sample:
            text = reader.pages[page_index].extract_text() or ""
            chars.append(sum(not char.isspace() for char in text))
        probe = "pdf_text_sample" if any(chars) else "pdf_no_text_in_sample"
        return probe, f"pages={count};sample_pages={','.join(str(i + 1) for i in sample)};nonspace_chars={','.join(map(str, chars))}"
    except Exception as error:  # One bad source should not stop the inventory.
        return "pdf_probe_error", f"{type(error).__name__}: {str(error)[:160]}"


def probe_txt(path: Path) -> tuple[str, str]:
    sample = path.open("rb").read(65536)
    if not sample:
        return "txt_empty", "bytes_sampled=0"
    if b"\x00" in sample:
        return "txt_binary_candidate", f"bytes_sampled={len(sample)}"
    for encoding in ("utf-8-sig", "gb18030"):
        try:
            sample.decode(encoding, errors="strict")
            return "txt_decodable_sample", f"encoding_candidate={encoding};bytes_sampled={len(sample)}"
        except UnicodeDecodeError:
            continue
    return "txt_encoding_review", f"bytes_sampled={len(sample)}"


def probe_docx(path: Path) -> tuple[str, str]:
    try:
        from docx import Document

        document = Document(str(path))
        chars = sum(len(p.text.strip()) for p in document.paragraphs[:50])
        return "docx_readable", f"paragraphs={len(document.paragraphs)};first_50_paragraph_chars={chars}"
    except Exception as error:
        return "docx_probe_error", f"{type(error).__name__}: {str(error)[:160]}"


def probe_signature(path: Path) -> tuple[str, str]:
    head = path.open("rb").read(16)
    if head.startswith(bytes.fromhex("D0CF11E0A1B11AE1")):
        return "legacy_ole_document", "binary Office format; text conversion pending"
    if head.startswith(b"{\\rtf"):
        return "rtf_document", "RTF signature"
    if head.startswith(b"PK"):
        return "zip_container", "ZIP signature"
    if head.startswith(b"Rar!"):
        return "rar_archive", "archive members not inventoried"
    return "unknown_signature", head.hex()


def probe(path: Path) -> tuple[str, str]:
    extension = path.suffix.lower()
    if extension == ".pdf":
        return probe_pdf(path)
    if extension == ".txt":
        return probe_txt(path)
    if extension == ".docx":
        return probe_docx(path)
    return probe_signature(path)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--source", type=Path, default=Path(r"F:\道教\六爻"))
    parser.add_argument("--output", type=Path, default=Path(__file__).resolve().parents[1] / "docs")
    args = parser.parse_args()
    source = args.source.resolve()
    output = args.output.resolve()
    if not source.is_dir():
        raise SystemExit(f"Source directory not found: {source}")
    output.mkdir(parents=True, exist_ok=True)

    files = sorted((p for p in source.rglob("*") if p.is_file()), key=lambda p: str(p.relative_to(source)).casefold())
    rows: list[dict[str, str | int]] = []
    by_hash: dict[str, str] = {}
    duplicates: dict[str, list[str]] = defaultdict(list)
    for path in files:
        relative = str(path.relative_to(source))
        digest = sha256_file(path)
        format_probe, sample_detail = probe(path)
        duplicate_of = by_hash.get(digest, "")
        if duplicate_of:
            duplicates[digest].append(relative)
        else:
            by_hash[digest] = relative
        rows.append({
            "relative_path": relative,
            "extension": path.suffix.lower(),
            "size_bytes": path.stat().st_size,
            "sha256": digest,
            "scope_candidate": scope_from_filename(path),
            "scope_basis": "filename_only",
            "format_probe": format_probe,
            "sample_detail": sample_detail,
            "duplicate_of": duplicate_of,
        })

    with (output / "source_inventory.csv").open("w", encoding="utf-8-sig", newline="") as stream:
        writer = csv.DictWriter(stream, fieldnames=FIELDS)
        writer.writeheader()
        writer.writerows(rows)
    summary = {
        "source": str(source),
        "file_count": len(rows),
        "total_bytes": sum(int(row["size_bytes"]) for row in rows),
        "extensions": dict(Counter(str(row["extension"]) for row in rows)),
        "scope_candidates": dict(Counter(str(row["scope_candidate"]) for row in rows)),
        "format_probes": dict(Counter(str(row["format_probe"]) for row in rows)),
        "exact_duplicate_groups": [
            [by_hash[digest], *members] for digest, members in duplicates.items()
        ],
    }
    (output / "source_audit.json").write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(summary, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
